import re
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.constants import OPEN_ASSIGNMENT
from app.deps import get_db, require_staff
from app.models import Assignment, Technician, User
from app.schemas import BulkResult, TechnicianIn, TechnicianOut
from app.utils import clean, commit_or_conflict, fold, ilike

router = APIRouter(prefix="/technicians", tags=["technicians"])


def to_technician(tech: Technician) -> TechnicianOut:
    return TechnicianOut(
        id=tech.id,
        name=tech.name,
        cargo=tech.cargo,
        dni=tech.dni,
        phone=tech.phone,
        archived=tech.archived,
        username=tech.user.username if tech.user else None,
    )


def load_technician(session: Session, technician_id: int) -> Technician:
    tech = session.scalar(
        select(Technician)
        .where(Technician.id == technician_id)
        .options(selectinload(Technician.user))
    )
    if tech is None:
        raise HTTPException(status_code=404, detail="No se encontró el técnico.")
    return tech


def person_name(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def person_key(value: str) -> str:
    return person_name(value).casefold()


def apply_fields(tech: Technician, name: str, cargo: str | None, dni: str | None, phone: str | None) -> None:
    tech.name = name
    tech.cargo = cargo
    tech.dni = dni
    tech.phone = phone
    if tech.user is not None:
        tech.user.full_name = name


@router.get("", response_model=list[TechnicianOut])
def list_technicians(
    q: str = "",
    include_archived: bool = False,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    stmt = select(Technician).options(selectinload(Technician.user)).order_by(Technician.name)
    if not include_archived:
        stmt = stmt.where(Technician.archived.is_(False))
    query = q.strip()
    if query:
        stmt = stmt.join(User, Technician.user_id == User.id, isouter=True).where(
            or_(ilike(Technician.name, query), ilike(Technician.dni, query), ilike(User.username, query))
        )
    return [to_technician(tech) for tech in session.scalars(stmt).all()]


@router.post("", response_model=TechnicianOut, status_code=201)
def create_technician(
    payload: TechnicianIn,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    name = person_name(payload.name)
    if len(name) < 3:
        raise HTTPException(status_code=400, detail="El nombre del técnico es obligatorio.")
    if session.scalar(select(Technician.id).where(func.lower(Technician.name) == name.lower())):
        raise HTTPException(status_code=409, detail="Ese técnico ya está en la lista.")
    tech = Technician(name=name, cargo=clean(payload.cargo), dni=clean(payload.dni), phone=clean(payload.phone))
    session.add(tech)
    commit_or_conflict(session, "Ese técnico ya está en la lista.")
    return to_technician(load_technician(session, tech.id))


@router.patch("/{technician_id}", response_model=TechnicianOut)
def update_technician(
    technician_id: int,
    payload: TechnicianIn,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    tech = load_technician(session, technician_id)
    name = person_name(payload.name)
    if len(name) < 3:
        raise HTTPException(status_code=400, detail="El nombre del técnico es obligatorio.")
    taken = session.scalar(
        select(Technician.id).where(func.lower(Technician.name) == name.lower(), Technician.id != tech.id)
    )
    if taken:
        raise HTTPException(status_code=409, detail="Ese técnico ya está en la lista.")
    apply_fields(tech, name, clean(payload.cargo), clean(payload.dni), clean(payload.phone))
    commit_or_conflict(session, "Ese técnico ya está en la lista.")
    return to_technician(load_technician(session, tech.id))


@router.post("/{technician_id}/archive", status_code=204)
def archive_technician(
    technician_id: int,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    tech = load_technician(session, technician_id)
    open_jobs = session.scalar(
        select(func.count())
        .select_from(Assignment)
        .where(Assignment.technician_id == tech.id, Assignment.status.in_(OPEN_ASSIGNMENT))
    )
    if open_jobs:
        raise HTTPException(
            status_code=400,
            detail="Este técnico tiene máquinas asignadas. Cancela la asignación o espera a que termine.",
        )
    tech.archived = True
    if tech.user:
        tech.user.active = False
    session.commit()


@router.post("/{technician_id}/restore", status_code=204)
def restore_technician(
    technician_id: int,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    tech = load_technician(session, technician_id)
    tech.archived = False
    if tech.user:
        tech.user.active = True
    session.commit()


def workbook_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Tecnicos"
    sheet.append(["tecnico", "cargo", "dni", "numero"])
    sheet.append(["PEREZ JUAN", "Técnico de reefer", "", ""])
    sheet.column_dimensions["A"].width = 36
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 16
    sheet.column_dimensions["D"].width = 16
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@router.get("/template")
def technicians_template(_: User = Depends(require_staff)):
    return Response(
        content=workbook_bytes(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla-tecnicos.xlsx"},
    )


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def header_key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", fold(cell_text(value)).lower())


@router.post("/bulk", response_model=BulkResult)
def bulk_technicians(
    file: UploadFile = File(...),
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Sube un Excel .xlsx. Puedes usar la plantilla.")
    raw = file.file.read(5_000_001)
    if len(raw) > 5_000_000:
        raise HTTPException(status_code=400, detail="El archivo supera 5 MB.")
    try:
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No se pudo leer el Excel.") from exc
    sheet = book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="El archivo está vacío.") from exc
    mapped = {header_key(name): index for index, name in enumerate(header or []) if header_key(name)}
    if "tecnico" not in mapped and "nombre" not in mapped:
        raise HTTPException(status_code=400, detail="La plantilla debe incluir la columna tecnico.")
    name_index = mapped.get("tecnico", mapped.get("nombre"))
    cargo_index = mapped.get("cargo")
    dni_index = mapped.get("dni")
    phone_index = mapped.get("numero", mapped.get("telefono"))

    existing = {
        person_key(tech.name): tech
        for tech in session.scalars(select(Technician).options(selectinload(Technician.user))).all()
    }
    created = 0
    updated = 0
    errors = []
    seen = 0
    for offset, row in enumerate(rows, start=2):
        if seen >= 5000:
            errors.append({"row": offset, "detail": "Se detuvo en 5000 filas."})
            break
        values = list(row or [])

        def take(index: int | None) -> str:
            if index is None or index >= len(values):
                return ""
            return cell_text(values[index])

        name = person_name(take(name_index))
        cargo = clean(take(cargo_index))
        dni = clean(take(dni_index))
        phone = clean(take(phone_index))
        if not name and not cargo and not dni and not phone:
            continue
        seen += 1
        if len(name) < 3:
            errors.append({"row": offset, "detail": "El nombre del técnico es obligatorio."})
            continue
        if cargo and len(cargo) > 120:
            errors.append({"row": offset, "detail": "El cargo es demasiado largo."})
            continue
        if dni and len(dni) > 20:
            errors.append({"row": offset, "detail": "El DNI es demasiado largo."})
            continue
        if phone and len(phone) > 30:
            errors.append({"row": offset, "detail": "El número es demasiado largo."})
            continue
        current = existing.get(person_key(name))
        if current is None:
            current = Technician(name=name, cargo=cargo, dni=dni, phone=phone)
            session.add(current)
            existing[person_key(name)] = current
            created += 1
            continue
        if current.archived:
            errors.append({"row": offset, "detail": "Está archivado. Puedes restaurarlo desde Auditoría."})
            continue
        changed = False
        if cargo and cargo != current.cargo:
            current.cargo = cargo
            changed = True
        if dni and dni != current.dni:
            current.dni = dni
            changed = True
        if phone and phone != current.phone:
            current.phone = phone
            changed = True
        if current.name != name:
            current.name = name
            if current.user is not None:
                current.user.full_name = name
            changed = True
        if changed:
            updated += 1
    book.close()
    commit_or_conflict(session, "No se pudo guardar la lista.")
    return BulkResult(created=created, updated=updated, errors=errors)
