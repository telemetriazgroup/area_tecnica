import csv
import io
import re
from io import BytesIO

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import FileResponse, Response
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session, selectinload

from app.constants import OPEN_ASSIGNMENT, STAFF
from app.deps import get_current_user, get_db, require_staff
from app.models import Assignment, Brand, Machine, ReeferModel, SparePart, User
from app.photos import photo_path, remove_file, save_image
from app.schemas import (
    AssignmentBrief,
    BrandOut,
    BulkResult,
    MachineIn,
    MachineOut,
    ModelOut,
    NameIn,
    SparePartIn,
    SparePartOut,
    SparePartUpdate,
)
from app.utils import clean, commit_or_conflict, fold, ilike, name_taken, normalize_code, normalize_name, normalize_serial

router = APIRouter(tags=["catalogo"])

def part_text_key(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip().casefold()


def description_taken(session: Session, model_id: int, description: str, exclude_id: int | None = None) -> bool:
    stmt = select(SparePart.id).where(
        SparePart.model_id == model_id,
        func.lower(SparePart.description) == description.lower(),
    )
    if exclude_id is not None:
        stmt = stmt.where(SparePart.id != exclude_id)
    return session.scalar(stmt) is not None


def to_model(model: ReeferModel) -> ModelOut:
    return ModelOut(id=model.id, brand_id=model.brand_id, name=model.name, archived=model.archived)


def to_brand(brand: Brand, include_archived: bool) -> BrandOut:
    models = [model for model in brand.models if include_archived or not model.archived]
    models.sort(key=lambda model: model.name.lower())
    return BrandOut(
        id=brand.id,
        name=brand.name,
        archived=brand.archived,
        models=[to_model(model) for model in models],
    )


def to_part(part: SparePart) -> SparePartOut:
    return SparePartOut(
        id=part.id,
        brand_id=part.brand_id,
        model_id=part.model_id,
        code=part.code,
        description=part.description,
        part_type=part.part_type,
        archived=part.archived,
    )


def get_brand(session: Session, brand_id: int) -> Brand:
    brand = session.scalar(select(Brand).where(Brand.id == brand_id).options(selectinload(Brand.models)))
    if brand is None:
        raise HTTPException(status_code=404, detail="No se encontró la marca.")
    return brand


def require_model(session: Session, brand_id: int, model_id: int, allow_archived: bool = False) -> ReeferModel:
    model = session.get(ReeferModel, model_id)
    if model is None or model.brand_id != brand_id:
        raise HTTPException(status_code=400, detail="El modelo no corresponde a la marca.")
    if model.archived and not allow_archived:
        raise HTTPException(status_code=400, detail="Ese modelo está archivado.")
    brand = session.get(Brand, brand_id)
    if brand is None:
        raise HTTPException(status_code=400, detail="No se encontró la marca.")
    if brand.archived and not allow_archived:
        raise HTTPException(status_code=400, detail="Esa marca está archivada.")
    return model


def to_machine(machine: Machine, assignment: Assignment | None) -> MachineOut:
    active = None
    if assignment is not None:
        active = AssignmentBrief(
            id=assignment.id,
            technician_id=assignment.technician_id,
            technician_name=assignment.technician.name,
            status=assignment.status,
        )
    return MachineOut(
        id=machine.id,
        serial=machine.serial,
        brand_id=machine.brand_id,
        brand_name=machine.brand.name,
        model_id=machine.model_id,
        model_name=machine.model.name,
        location=machine.location,
        year=machine.year,
        archived=machine.archived,
        has_machine_photo=bool(machine.machine_photo),
        has_serial_photo=bool(machine.serial_photo),
        active_assignment=active,
    )


@router.get("/brands", response_model=list[BrandOut])
def list_brands(
    include_archived: bool = False,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    brands = session.scalars(select(Brand).options(selectinload(Brand.models)).order_by(Brand.name)).all()
    if not include_archived:
        brands = [brand for brand in brands if not brand.archived]
    return [to_brand(brand, include_archived) for brand in brands]


@router.post("/brands", response_model=BrandOut, status_code=201)
def create_brand(payload: NameIn, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    name = normalize_name(payload.name, "La marca")
    if name_taken(session, Brand, name):
        raise HTTPException(status_code=409, detail="Esa marca ya existe.")
    brand = Brand(name=name)
    session.add(brand)
    commit_or_conflict(session, "Esa marca ya existe.")
    session.refresh(brand)
    brand.models = []
    return to_brand(brand, True)


@router.patch("/brands/{brand_id}", response_model=BrandOut)
def update_brand(
    brand_id: int,
    payload: NameIn,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    brand = get_brand(session, brand_id)
    name = normalize_name(payload.name, "La marca")
    if name_taken(session, Brand, name, exclude_id=brand.id):
        raise HTTPException(status_code=409, detail="Esa marca ya existe.")
    brand.name = name
    commit_or_conflict(session, "Esa marca ya existe.")
    return to_brand(get_brand(session, brand.id), True)


@router.post("/brands/{brand_id}/archive", status_code=204)
def archive_brand(brand_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    brand = get_brand(session, brand_id)
    brand.archived = True
    session.commit()


@router.post("/brands/{brand_id}/restore", status_code=204)
def restore_brand(brand_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    brand = get_brand(session, brand_id)
    brand.archived = False
    session.commit()


def catalog_label(value: str) -> str:
    return re.sub(r"\s+", " ", value or "").strip()


def catalog_key(value: str) -> str:
    return catalog_label(value).casefold()


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def header_key(value) -> str:
    return re.sub(r"[^a-z0-9]", "", fold(cell_text(value)).lower())


def brand_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Marcas"
    sheet.append(["MARCA", "MODELO"])
    sheet.append(["THERMO KING", "MP 5000"])
    sheet.append(["DAIKIN", ""])
    sheet.column_dimensions["A"].width = 24
    sheet.column_dimensions["B"].width = 24
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@router.get("/brands/template")
def brands_template(_: User = Depends(require_staff)):
    return Response(
        content=brand_workbook(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla-marcas.xlsx"},
    )


def model_on(brand: Brand, name: str) -> ReeferModel | None:
    wanted = catalog_key(name)
    return next((model for model in brand.models if catalog_key(model.name) == wanted), None)


@router.post("/brands/bulk", response_model=BulkResult)
def bulk_brands(
    file: UploadFile = File(...),
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    filename = (file.filename or "").lower()
    if not filename.endswith(".xlsx"):
        raise HTTPException(status_code=400, detail="Sube un Excel .xlsx. Puedes usar la plantilla.")
    raw = file.file.read(5_000_001)
    if len(raw) > 5_000_000:
        raise HTTPException(status_code=400, detail="El archivo supera 5 MB.")
    try:
        book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
    except Exception as exc:
        raise HTTPException(status_code=400, detail="No se pudo leer el Excel.") from exc
    sheet = book.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    try:
        header = next(rows)
    except StopIteration as exc:
        raise HTTPException(status_code=400, detail="El archivo está vacío.") from exc
    mapped = {header_key(name): index for index, name in enumerate(header or []) if header_key(name)}
    if "marca" not in mapped:
        raise HTTPException(status_code=400, detail="La plantilla debe incluir la columna MARCA.")
    brand_index = mapped["marca"]
    model_index = mapped.get("modelo")
    brands = {
        catalog_key(brand.name): brand
        for brand in session.scalars(select(Brand).options(selectinload(Brand.models))).all()
    }
    created = 0
    updated = 0
    errors = []
    seen = 0
    for offset, row in enumerate(rows, start=2):
        if seen >= 5000:
            errors.append({"row": offset, "detail": "Se detuvo en 5000 filas."})
            break
        values = list(row or [])

        def take(index: int | None) -> str:
            if index is None or index >= len(values):
                return ""
            return cell_text(values[index])

        brand_name = catalog_label(take(brand_index))
        model_name = catalog_label(take(model_index))
        if not brand_name and not model_name:
            continue
        seen += 1
        if len(brand_name) < 2:
            errors.append({"row": offset, "detail": "La marca es obligatoria."})
            continue
        if len(brand_name) > 120 or len(model_name) > 120:
            errors.append({"row": offset, "detail": "El nombre es demasiado largo."})
            continue
        if model_name and len(model_name) < 2:
            errors.append({"row": offset, "detail": "El modelo es demasiado corto."})
            continue
        brand = brands.get(catalog_key(brand_name))
        if brand is None:
            brand = Brand(name=brand_name)
            session.add(brand)
            brand.models = []
            brands[catalog_key(brand_name)] = brand
            created += 1
        elif brand.archived:
            errors.append({"row": offset, "detail": "La marca está archivada. Puedes restaurarla desde Auditoría."})
            continue
        elif brand.name != brand_name:
            brand.name = brand_name
            updated += 1
        if not model_name:
            continue
        model = model_on(brand, model_name)
        if model is None:
            model = ReeferModel(name=model_name)
            brand.models.append(model)
            session.add(model)
            created += 1
            continue
        if model.archived:
            errors.append({"row": offset, "detail": "El modelo está archivado. Puedes restaurarlo desde Auditoría."})
            continue
        if model.name != model_name:
            model.name = model_name
            updated += 1
    book.close()
    commit_or_conflict(session, "No se pudo guardar las marcas.")
    return BulkResult(created=created, updated=updated, errors=errors)


@router.post("/brands/{brand_id}/models", response_model=ModelOut, status_code=201)
def create_model(
    brand_id: int,
    payload: NameIn,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    brand = get_brand(session, brand_id)
    if brand.archived:
        raise HTTPException(status_code=400, detail="La marca está archivada.")
    name = normalize_name(payload.name, "El modelo")
    if name_taken(session, ReeferModel, name, brand_id=brand.id):
        raise HTTPException(status_code=409, detail="Ese modelo ya existe en la marca.")
    model = ReeferModel(brand_id=brand.id, name=name)
    session.add(model)
    commit_or_conflict(session, "Ese modelo ya existe en la marca.")
    session.refresh(model)
    return to_model(model)


@router.patch("/models/{model_id}", response_model=ModelOut)
def update_model(
    model_id: int,
    payload: NameIn,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    model = session.get(ReeferModel, model_id)
    if model is None:
        raise HTTPException(status_code=404, detail="No se encontró el modelo.")
    name = normalize_name(payload.name, "El modelo")
    if name_taken(session, ReeferModel, name, brand_id=model.brand_id, exclude_id=model.id):
        raise HTTPException(status_code=409, detail="Ese modelo ya existe en la marca.")
    model.name = name
    commit_or_conflict(session, "Ese modelo ya existe en la marca.")
    session.refresh(model)
    return to_model(model)


@router.post("/models/{model_id}/archive", status_code=204)
def archive_model(model_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    model = session.get(ReeferModel, model_id)
    if model is None:
        raise HTTPException(status_code=404, detail="No se encontró el modelo.")
    model.archived = True
    session.commit()


@router.post("/models/{model_id}/restore", status_code=204)
def restore_model(model_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    model = session.get(ReeferModel, model_id)
    if model is None:
        raise HTTPException(status_code=404, detail="No se encontró el modelo.")
    model.archived = False
    session.commit()


@router.get("/machines", response_model=list[MachineOut])
def list_machines(
    q: str = "",
    include_archived: bool = False,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    stmt = (
        select(Machine)
        .options(selectinload(Machine.brand), selectinload(Machine.model))
        .order_by(Machine.serial)
    )
    if not include_archived:
        stmt = stmt.where(Machine.archived.is_(False))
    query = q.strip()
    if query:
        stmt = stmt.where(or_(ilike(Machine.serial, query), ilike(Machine.location, query)))
    machines = session.scalars(stmt).all()
    ids = [machine.id for machine in machines]
    assignments = []
    if ids:
        assignments = session.scalars(
            select(Assignment)
            .where(Assignment.machine_id.in_(ids), Assignment.status.in_(OPEN_ASSIGNMENT))
            .options(selectinload(Assignment.technician))
        ).all()
    by_machine = {assignment.machine_id: assignment for assignment in assignments}
    return [to_machine(machine, by_machine.get(machine.id)) for machine in machines]


@router.post("/machines", response_model=MachineOut, status_code=201)
def create_machine(payload: MachineIn, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    serial = normalize_serial(payload.serial)
    model = require_model(session, payload.brand_id, payload.model_id)
    if session.scalar(select(Machine.id).where(Machine.serial == serial)):
        raise HTTPException(status_code=409, detail="Esa serie ya está registrada.")
    machine = Machine(
        serial=serial,
        brand_id=payload.brand_id,
        model_id=model.id,
        location=clean(payload.location),
        year=payload.year,
    )
    session.add(machine)
    commit_or_conflict(session, "Esa serie ya está registrada.")
    session.refresh(machine)
    machine.brand = session.get(Brand, machine.brand_id)
    machine.model = model
    return to_machine(machine, None)


@router.patch("/machines/{machine_id}", response_model=MachineOut)
def update_machine(
    machine_id: int,
    payload: MachineIn,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    machine = session.scalar(
        select(Machine)
        .where(Machine.id == machine_id)
        .options(selectinload(Machine.brand), selectinload(Machine.model))
    )
    if machine is None:
        raise HTTPException(status_code=404, detail="No se encontró la máquina.")
    serial = normalize_serial(payload.serial)
    model = require_model(session, payload.brand_id, payload.model_id, allow_archived=True)
    taken = session.scalar(select(Machine.id).where(Machine.serial == serial, Machine.id != machine.id))
    if taken:
        raise HTTPException(status_code=409, detail="Esa serie ya está registrada.")
    machine.serial = serial
    machine.brand_id = payload.brand_id
    machine.model_id = model.id
    machine.location = clean(payload.location)
    machine.year = payload.year
    commit_or_conflict(session, "Esa serie ya está registrada.")
    machine = session.scalar(
        select(Machine)
        .where(Machine.id == machine_id)
        .options(selectinload(Machine.brand), selectinload(Machine.model))
    )
    assignment = session.scalar(
        select(Assignment)
        .where(Assignment.machine_id == machine_id, Assignment.status.in_(OPEN_ASSIGNMENT))
        .options(selectinload(Assignment.technician))
    )
    return to_machine(machine, assignment)


def _machine_or_404(session: Session, machine_id: int) -> Machine:
    machine = session.get(Machine, machine_id)
    if machine is None:
        raise HTTPException(status_code=404, detail="No se encontró la máquina.")
    return machine


@router.post("/machines/{machine_id}/archive", status_code=204)
def archive_machine(machine_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    machine = _machine_or_404(session, machine_id)
    open_job = session.scalar(
        select(Assignment.id).where(Assignment.machine_id == machine.id, Assignment.status.in_(OPEN_ASSIGNMENT))
    )
    if open_job:
        raise HTTPException(
            status_code=400,
            detail="Esta máquina tiene una revisión abierta. Cancélala o espera a que termine.",
        )
    machine.archived = True
    session.commit()


@router.post("/machines/{machine_id}/restore", status_code=204)
def restore_machine(machine_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    machine = _machine_or_404(session, machine_id)
    machine.archived = False
    session.commit()


PHOTO_FIELDS = {"maquina": "machine_photo", "serie": "serial_photo"}


def photo_field(kind: str) -> str:
    field = PHOTO_FIELDS.get(kind)
    if field is None:
        raise HTTPException(status_code=404, detail="No se encontró la foto.")
    return field


def can_view_machine(session: Session, user: User, machine_id: int) -> None:
    if user.role in STAFF:
        return
    technician_id = user.technician.id if user.technician else None
    if technician_id is None:
        raise HTTPException(status_code=403, detail="No tienes permiso para ver esta foto.")
    assigned = session.scalar(
        select(Assignment.id).where(
            Assignment.machine_id == machine_id,
            Assignment.technician_id == technician_id,
        )
    )
    if assigned is None:
        raise HTTPException(status_code=403, detail="No tienes permiso para ver esta foto.")


@router.get("/machines/{machine_id}/photos/{kind}")
def machine_photo(
    machine_id: int,
    kind: str,
    session: Session = Depends(get_db),
    user: User = Depends(get_current_user),
):
    machine = _machine_or_404(session, machine_id)
    can_view_machine(session, user, machine.id)
    filename = getattr(machine, photo_field(kind))
    if not filename:
        raise HTTPException(status_code=404, detail="Esta máquina no tiene esa foto.")
    return FileResponse(photo_path(filename), media_type="image/jpeg")


@router.post("/machines/{machine_id}/photos/{kind}", response_model=MachineOut)
def upload_machine_photo(
    machine_id: int,
    kind: str,
    photo: UploadFile = File(...),
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    field = photo_field(kind)
    machine = session.scalar(
        select(Machine)
        .where(Machine.id == machine_id)
        .options(selectinload(Machine.brand), selectinload(Machine.model))
    )
    if machine is None:
        raise HTTPException(status_code=404, detail="No se encontró la máquina.")
    if not photo.filename:
        raise HTTPException(status_code=400, detail="La foto está vacía.")
    filename = save_image(photo)
    previous = getattr(machine, field)
    setattr(machine, field, filename)
    session.commit()
    if previous:
        remove_file(previous)
    assignment = session.scalar(
        select(Assignment)
        .where(Assignment.machine_id == machine.id, Assignment.status.in_(OPEN_ASSIGNMENT))
        .options(selectinload(Assignment.technician))
    )
    return to_machine(machine, assignment)


@router.delete("/machines/{machine_id}/photos/{kind}", response_model=MachineOut)
def delete_machine_photo(
    machine_id: int,
    kind: str,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    field = photo_field(kind)
    machine = session.scalar(
        select(Machine)
        .where(Machine.id == machine_id)
        .options(selectinload(Machine.brand), selectinload(Machine.model))
    )
    if machine is None:
        raise HTTPException(status_code=404, detail="No se encontró la máquina.")
    previous = getattr(machine, field)
    setattr(machine, field, None)
    session.commit()
    if previous:
        remove_file(previous)
    assignment = session.scalar(
        select(Assignment)
        .where(Assignment.machine_id == machine.id, Assignment.status.in_(OPEN_ASSIGNMENT))
        .options(selectinload(Assignment.technician))
    )
    return to_machine(machine, assignment)


@router.get("/spare-parts", response_model=list[SparePartOut])
def list_parts(
    model_id: int | None = None,
    q: str = "",
    include_archived: bool = False,
    limit: int = 50,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    if model_id is None:
        return []
    limit = min(max(limit, 1), 5000)
    stmt = select(SparePart).where(SparePart.model_id == model_id).order_by(SparePart.description).limit(limit)
    if not include_archived:
        stmt = stmt.where(SparePart.archived.is_(False))
    query = q.strip()
    if query:
        stmt = stmt.where(or_(ilike(SparePart.description, query), ilike(SparePart.code, query)))
    return [to_part(part) for part in session.scalars(stmt).all()]


@router.post("/spare-parts", response_model=SparePartOut, status_code=201)
def create_part(payload: SparePartIn, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    model = require_model(session, payload.brand_id, payload.model_id)
    code = normalize_code(payload.code)
    description = normalize_name(payload.description, "La descripción", 200)
    existing = session.scalar(
        select(SparePart).where(SparePart.model_id == model.id, func.lower(SparePart.code) == code.lower())
    )
    if existing:
        raise HTTPException(status_code=409, detail="Esa serie ya está en el catálogo de este modelo.")
    if description_taken(session, model.id, description):
        raise HTTPException(status_code=409, detail="Esa descripción ya está en el catálogo de este modelo.")
    part = SparePart(
        brand_id=payload.brand_id,
        model_id=model.id,
        code=code,
        description=description,
        part_type=clean(payload.part_type),
    )
    session.add(part)
    commit_or_conflict(session, "Esa serie ya está en el catálogo de este modelo.")
    session.refresh(part)
    return to_part(part)


def parts_workbook() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Repuestos"
    sheet.append(["serie", "descripcion", "tipo"])
    sheet.append(["818521B", "CONTACTOR", ""])
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 42
    sheet.column_dimensions["C"].width = 18
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


@router.get("/spare-parts/template")
def parts_template(_: User = Depends(require_staff)):
    return Response(
        content=parts_workbook(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=plantilla-repuestos.xlsx"},
    )


def sheet_rows(raw: bytes, filename: str) -> list[tuple[int, str, str, str]]:
    if filename.endswith(".xlsx"):
        try:
            book = load_workbook(BytesIO(raw), read_only=True, data_only=True)
        except Exception as exc:
            raise HTTPException(status_code=400, detail="No se pudo leer el Excel.") from exc
        sheet = book.worksheets[0]
        rows = sheet.iter_rows(values_only=True)
        try:
            header = next(rows)
        except StopIteration as exc:
            raise HTTPException(status_code=400, detail="El archivo está vacío.") from exc
        mapped = {header_key(name): index for index, name in enumerate(header or []) if header_key(name)}
        if "serie" not in mapped and "codigo" not in mapped and "code" not in mapped:
            raise HTTPException(status_code=400, detail="La plantilla debe incluir las columnas serie y descripcion.")
        if "descripcion" not in mapped and "description" not in mapped:
            raise HTTPException(status_code=400, detail="La plantilla debe incluir las columnas serie y descripcion.")
        code_index = mapped.get("serie", mapped.get("codigo", mapped.get("code")))
        description_index = mapped.get("descripcion", mapped.get("description"))
        type_index = mapped.get("tipo", mapped.get("type"))
        parsed = []
        for offset, row in enumerate(rows, start=2):
            values = list(row or [])

            def take(index: int | None) -> str:
                if index is None or index >= len(values):
                    return ""
                return cell_text(values[index])

            parsed.append((offset, take(code_index), take(description_index), take(type_index)))
        book.close()
        return parsed
    if filename.endswith(".xls"):
        raise HTTPException(status_code=400, detail="Usa un Excel .xlsx. Puedes descargar la plantilla.")
    try:
        text = raw.decode("utf-8-sig")
    except UnicodeError as exc:
        raise HTTPException(status_code=400, detail="El archivo debe estar en UTF-8.") from exc
    if not text.strip():
        raise HTTPException(status_code=400, detail="El archivo está vacío.")
    first = text.splitlines()[0]
    delimiter = ";" if first.count(";") > first.count(",") else ","
    reader = csv.DictReader(io.StringIO(text), delimiter=delimiter)
    if not reader.fieldnames:
        raise HTTPException(status_code=400, detail="El archivo no tiene encabezados.")
    mapped = {header_key(name): name for name in reader.fieldnames if name}
    code_key = mapped.get("serie") or mapped.get("codigo") or mapped.get("code")
    desc_key = mapped.get("descripcion") or mapped.get("description")
    type_key = mapped.get("tipo") or mapped.get("type")
    if not code_key or not desc_key:
        raise HTTPException(status_code=400, detail="La plantilla debe incluir las columnas serie y descripcion.")
    parsed = []
    for offset, row in enumerate(reader, start=2):
        parsed.append((offset, (row.get(code_key) or "").strip(), (row.get(desc_key) or "").strip(), (row.get(type_key) or "").strip() if type_key else ""))
    return parsed


@router.post("/spare-parts/bulk", response_model=BulkResult)
def bulk_parts(
    brand_id: int,
    model_id: int,
    file: UploadFile = File(...),
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    model = require_model(session, brand_id, model_id)
    filename = (file.filename or "").lower()
    raw = file.file.read(5_000_001)
    if len(raw) > 5_000_000:
        raise HTTPException(status_code=400, detail="El archivo supera 5 MB.")
    parsed = sheet_rows(raw, filename)
    current = session.scalars(select(SparePart).where(SparePart.model_id == model.id)).all()
    by_code = {part.code.casefold(): part for part in current}
    by_description = {part_text_key(part.description): part for part in current}
    seen_codes: set[str] = set()
    seen_descriptions: set[str] = set()
    created = 0
    updated = 0
    errors = []
    seen = 0
    for index, code_raw, description_raw, type_raw in parsed:
        if seen >= 5000:
            errors.append({"row": index, "detail": "Se detuvo en 5000 filas."})
            break
        if not code_raw and not description_raw and not type_raw:
            continue
        seen += 1
        try:
            code = normalize_code(code_raw)
            description = normalize_name(description_raw, "La descripción", 200)
        except HTTPException as exc:
            errors.append({"row": index, "detail": str(exc.detail)})
            continue
        code_key = code.casefold()
        desc_key = part_text_key(description)
        if code_key in seen_codes:
            errors.append({"row": index, "detail": "La serie está repetida en el archivo."})
            continue
        if desc_key in seen_descriptions:
            errors.append({"row": index, "detail": "La descripción está repetida en el archivo."})
            continue
        seen_codes.add(code_key)
        seen_descriptions.add(desc_key)
        part = by_code.get(code_key)
        described = by_description.get(desc_key)
        if part is None:
            if described is not None:
                errors.append({"row": index, "detail": "Esa descripción ya está en el catálogo de este modelo."})
                continue
            part = SparePart(
                brand_id=brand_id,
                model_id=model.id,
                code=code,
                description=description,
                part_type=clean(type_raw),
            )
            session.add(part)
            by_code[code_key] = part
            by_description[desc_key] = part
            created += 1
            continue
        if part.archived:
            errors.append({"row": index, "detail": "Está archivado. Puedes restaurarlo desde Auditoría."})
            continue
        if described is not None and described is not part:
            errors.append({"row": index, "detail": "Esa descripción ya está en el catálogo de este modelo."})
            continue
        changed = False
        if part.description != description:
            by_description.pop(part_text_key(part.description), None)
            part.description = description
            by_description[desc_key] = part
            changed = True
        part_type = clean(type_raw)
        if part_type and part_type != part.part_type:
            part.part_type = part_type
            changed = True
        if changed:
            updated += 1
    commit_or_conflict(session, "No se pudo guardar la lista. Revisa series o descripciones duplicadas.")
    return BulkResult(created=created, updated=updated, errors=errors)


@router.patch("/spare-parts/{part_id}", response_model=SparePartOut)
def update_part(
    part_id: int,
    payload: SparePartUpdate,
    session: Session = Depends(get_db),
    _: User = Depends(require_staff),
):
    part = session.get(SparePart, part_id)
    if part is None:
        raise HTTPException(status_code=404, detail="No se encontró el repuesto.")
    code = normalize_code(payload.code)
    description = normalize_name(payload.description, "La descripción", 200)
    taken = session.scalar(
        select(SparePart.id).where(
            SparePart.model_id == part.model_id,
            func.lower(SparePart.code) == code.lower(),
            SparePart.id != part.id,
        )
    )
    if taken:
        raise HTTPException(status_code=409, detail="Esa serie ya está en el catálogo de este modelo.")
    if description_taken(session, part.model_id, description, exclude_id=part.id):
        raise HTTPException(status_code=409, detail="Esa descripción ya está en el catálogo de este modelo.")
    part.code = code
    part.description = description
    part.part_type = clean(payload.part_type)
    commit_or_conflict(session, "Esa serie ya está en el catálogo de este modelo.")
    session.refresh(part)
    return to_part(part)


@router.post("/spare-parts/{part_id}/archive", status_code=204)
def archive_part(part_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    part = session.get(SparePart, part_id)
    if part is None:
        raise HTTPException(status_code=404, detail="No se encontró el repuesto.")
    part.archived = True
    session.commit()


@router.post("/spare-parts/{part_id}/restore", status_code=204)
def restore_part(part_id: int, session: Session = Depends(get_db), _: User = Depends(require_staff)):
    part = session.get(SparePart, part_id)
    if part is None:
        raise HTTPException(status_code=404, detail="No se encontró el repuesto.")
    part.archived = False
    session.commit()
