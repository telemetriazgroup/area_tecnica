from datetime import date, datetime
from io import BytesIO
from pathlib import Path
from zoneinfo import ZoneInfo

from fpdf import FPDF
from openpyxl import Workbook
from openpyxl.drawing.image import Image as SheetImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.config import settings
from app.schemas import DailyReport, NeedsReport, PartSum

LOGO = Path(__file__).resolve().parent / "assets" / "logo.png"
NAVY = (24, 44, 77)
NAVY_HEX = "182C4D"
MUTED = (90, 106, 128)
INK = (21, 32, 51)
ZEBRA = (232, 237, 244)
CONDITIONS = {
    "faltante": "Faltante",
    "incompleto": "Incompleto",
    "inoperativo": "Inoperativo",
}

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PDF_MEDIA = "application/pdf"


def format_day(value: date) -> str:
    return value.strftime("%d/%m/%Y")


def format_when(value: datetime | None) -> str:
    if value is None:
        return ""
    local = value.astimezone(ZoneInfo(settings.timezone))
    return local.strftime("%d/%m/%Y %H:%M")


def format_duration(start: datetime | None, end: datetime | None) -> str:
    if start is None or end is None:
        return ""
    seconds = max(0, int((end - start).total_seconds()))
    days, seconds = divmod(seconds, 86400)
    hours, seconds = divmod(seconds, 3600)
    minutes = seconds // 60
    if days:
        return f"{days} d {hours} h"
    if hours:
        return f"{hours} h {minutes} min"
    if minutes:
        return f"{minutes} min"
    return "menos de 1 min"


def pdf_text(value) -> str:
    return str(value or "").replace("\n", " ").encode("latin-1", "replace").decode("latin-1")


def place_label(location: str | None, year: int | None = None) -> str:
    parts = [location or "", str(year) if year else ""]
    return " · ".join(part for part in parts if part)


class ReportPDF(FPDF):
    def __init__(self, title: str, subtitle: str):
        super().__init__(orientation="L", unit="mm", format="A4")
        self.report_title = title
        self.report_subtitle = subtitle
        self.set_auto_page_break(auto=True, margin=14)
        self.set_margins(12, 36, 12)
        self.alias_nb_pages()

    def header(self):
        self.image(str(LOGO), x=12, y=8, w=58)
        self.set_xy(74, 11)
        self.set_text_color(*NAVY)
        self.set_font("Helvetica", "B", 16)
        self.cell(0, 8, pdf_text(self.report_title), new_x="LMARGIN", new_y="NEXT")
        self.set_x(74)
        self.set_text_color(*MUTED)
        self.set_font("Helvetica", "", 11)
        self.cell(0, 6, pdf_text(self.report_subtitle))
        self.set_draw_color(*NAVY)
        self.set_line_width(0.7)
        self.line(12, 30, self.w - 12, 30)
        self.set_y(self.t_margin)

    def footer(self):
        self.set_y(-11)
        self.set_font("Helvetica", "", 8)
        self.set_text_color(*MUTED)
        self.cell(0, 6, pdf_text(f"Área técnica  ·  Página {self.page_no()}/{{nb}}"), align="R")


def fit(pdf: FPDF, text, width: float) -> str:
    value = pdf_text(text)
    if pdf.get_string_width(value) <= width - 2:
        return value
    trimmed = value
    while trimmed and pdf.get_string_width(trimmed + "...") > width - 2:
        trimmed = trimmed[:-1]
    return f"{trimmed}..."


def paint_head(pdf: FPDF, headers: list[str], widths: list[float]) -> None:
    pdf.set_font("Helvetica", "B", 9)
    pdf.set_fill_color(*NAVY)
    pdf.set_text_color(255, 255, 255)
    for header, width in zip(headers, widths):
        pdf.cell(width, 8, pdf_text(header), fill=True)
    pdf.ln()
    pdf.set_text_color(*INK)
    pdf.set_font("Helvetica", "", 9)


def paint_rows(pdf: FPDF, rows: list[list], widths: list[float], headers: list[str]) -> None:
    if not rows:
        pdf.set_font("Helvetica", "", 10)
        pdf.set_text_color(*MUTED)
        pdf.cell(0, 8, pdf_text("Sin registros en este corte."))
        pdf.ln(10)
        pdf.set_text_color(*INK)
        return
    for index, row in enumerate(rows):
        if pdf.get_y() + 8 > pdf.page_break_trigger:
            pdf.add_page()
            paint_head(pdf, headers, widths)
        pdf.set_fill_color(*(ZEBRA if index % 2 else (255, 255, 255)))
        for value, width in zip(row, widths):
            pdf.cell(width, 7, fit(pdf, value, width), fill=True)
        pdf.ln()
    pdf.ln(4)


def build_daily_pdf(report: DailyReport) -> bytes:
    pdf = ReportPDF("Máquinas inspeccionadas", format_day(report.date))
    pdf.add_page()
    headers = ["Serie", "Marca", "Modelo", "Ubicación", "Repuestos", "Tiempo", "Cierre"]
    widths = [40, 42, 42, 42, 26, 34, 47]
    if not report.groups:
        pdf.set_font("Helvetica", "", 11)
        pdf.set_text_color(*MUTED)
        pdf.cell(0, 8, pdf_text("Ese día no se cerró ninguna inspección."))
        return bytes(pdf.output())
    for group in report.groups:
        if pdf.get_y() + 24 > pdf.page_break_trigger:
            pdf.add_page()
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 8, pdf_text(group.technician), new_x="LMARGIN", new_y="NEXT")
        paint_head(pdf, headers, widths)
        paint_rows(
            pdf,
            [
                [
                    row.serial,
                    row.brand,
                    row.model,
                    row.location or "",
                    row.items_count,
                    format_duration(row.assigned_at, row.completed_at),
                    format_when(row.completed_at),
                ]
                for row in group.inspections
            ],
            widths,
            headers,
        )
    return bytes(pdf.output())


def build_needs_pdf(report: NeedsReport, view: str) -> bytes:
    if view == "general":
        return general_pdf(report)
    if view == "modelo":
        block = report.by_model
        pieces = sum(part.total for part in block.parts)
        title = f"Necesidad de {block.brand} {block.model}"
        subtitle = (
            f"Del {format_day(report.date_from)} al {format_day(report.date_to)}"
            f"  ·  {block.machines} máquinas  ·  {pieces} piezas"
        )
    elif view == "total":
        title = "Conglomerado de repuestos"
        subtitle = (
            f"Del {format_day(report.date_from)} al {format_day(report.date_to)}"
            f"  ·  {report.summary.inspections} inspecciones  ·  {report.summary.parts_total} piezas"
        )
    else:
        title = "Necesidad por serie"
        subtitle = (
            f"Del {format_day(report.date_from)} al {format_day(report.date_to)}"
            f"  ·  {report.summary.inspections} inspecciones  ·  {report.summary.parts_total} piezas"
        )
    pdf = ReportPDF(title, subtitle)
    pdf.add_page()
    if view == "serie":
        write_serial_blocks(pdf, report.by_serial, "No hay inspecciones cerradas en ese rango.")
    elif view == "modelo":
        write_serial_blocks(pdf, report.by_model.units, "Ese modelo no tiene inspecciones cerradas en ese rango.")
        if report.by_model.units:
            if pdf.get_y() + 20 > pdf.page_break_trigger:
                pdf.add_page()
            pdf.set_text_color(*NAVY)
            pdf.set_font("Helvetica", "B", 12)
            pdf.cell(0, 8, pdf_text("Total del modelo"), new_x="LMARGIN", new_y="NEXT")
            write_sum_pdf(pdf, report.by_model.parts)
    else:
        write_sum_pdf(pdf, report.conglomerate)
    return bytes(pdf.output())


def general_pdf(report: NeedsReport) -> bytes:
    span = f"Del {format_day(report.date_from)} al {format_day(report.date_to)}"
    pdf = ReportPDF(
        "Reporte general",
        f"{span}  ·  {report.summary.machines} máquinas  ·  {report.summary.parts_total} piezas",
    )
    if not report.models:
        pdf.add_page()
        pdf.set_font("Helvetica", "", 11)
        pdf.set_text_color(*MUTED)
        pdf.cell(0, 8, pdf_text("No hay inspecciones cerradas en ese rango."))
        return bytes(pdf.output())
    for block in report.models:
        pieces = sum(part.total for part in block.parts)
        pdf.report_title = f"{block.brand} {block.model}"
        pdf.report_subtitle = f"{span}  ·  {block.machines} máquinas  ·  {pieces} piezas"
        pdf.add_page()
        write_serial_blocks(pdf, block.units, "Sin equipos en este modelo.")
    return bytes(pdf.output())


def write_sum_pdf(pdf: FPDF, parts: list[PartSum]) -> None:
    headers = ["Serie", "Descripción", "Tipo", "Faltante", "Incompleto", "Inoperativo", "Total"]
    widths = [36, 100, 34, 26, 28, 30, 19]
    paint_head(pdf, headers, widths)
    paint_rows(
        pdf,
        [
            [part.code, part.description, part.part_type or "", part.faltante, part.incompleto, part.inoperativo, part.total]
            for part in parts
        ],
        widths,
        headers,
    )


def write_serial_blocks(pdf: FPDF, blocks: list, empty: str) -> None:
    headers = ["Cant.", "Condición", "Serie", "Descripción", "Nota"]
    widths = [18, 32, 36, 110, 77]
    if not blocks:
        pdf.set_font("Helvetica", "", 11)
        pdf.set_text_color(*MUTED)
        pdf.cell(0, 8, pdf_text(empty))
        return
    for block in blocks:
        detail = " · ".join(
            part
            for part in (
                f"{block.brand} {block.model}",
                place_label(block.location, block.year),
                block.technician,
                format_when(block.completed_at),
            )
            if part
        )
        if pdf.get_y() + 22 > pdf.page_break_trigger:
            pdf.add_page()
        pdf.set_text_color(*NAVY)
        pdf.set_font("Helvetica", "B", 12)
        pdf.cell(0, 7, pdf_text(block.serial), new_x="LMARGIN", new_y="NEXT")
        pdf.set_text_color(*MUTED)
        pdf.set_font("Helvetica", "", 9)
        pdf.cell(0, 5, pdf_text(detail), new_x="LMARGIN", new_y="NEXT")
        if not block.parts:
            pdf.set_text_color(*INK)
            pdf.cell(0, 6, pdf_text("Sin faltantes."), new_x="LMARGIN", new_y="NEXT")
            pdf.ln(2)
            continue
        paint_head(pdf, headers, widths)
        paint_rows(
            pdf,
            [
                [
                    part.quantity,
                    CONDITIONS.get(part.condition, part.condition),
                    part.code,
                    part.description,
                    part.note or "",
                ]
                for part in block.parts
            ],
            widths,
            headers,
        )


def style_header(cell) -> None:
    cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
    cell.fill = PatternFill("solid", fgColor=NAVY_HEX)
    cell.alignment = Alignment(vertical="center")


def write_table(sheet: Worksheet, row_index: int, headers: list[str], rows: list[list]) -> None:
    thin = Border(bottom=Side(style="thin", color="D5DEEA"))
    for column, header in enumerate(headers, start=1):
        cell = sheet.cell(row_index, column, header)
        style_header(cell)
    if not rows:
        sheet.cell(row_index + 1, 1, "Sin registros en este corte.")
        return
    for offset, values in enumerate(rows, start=1):
        fill = PatternFill("solid", fgColor="E8EEF6") if offset % 2 == 0 else PatternFill("solid", fgColor="FFFFFF")
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row_index + offset, column, value)
            cell.font = Font(name="Calibri", size=11, color="152033")
            cell.fill = fill
            cell.border = thin
            cell.alignment = Alignment(vertical="center")
    sheet.auto_filter.ref = f"A{row_index}:{get_column_letter(len(headers))}{row_index + len(rows)}"
    sheet.freeze_panes = f"A{row_index + 1}"
    for column, header in enumerate(headers, start=1):
        width = max(len(header) + 2, 12)
        for values in rows[:80]:
            width = max(width, min(48, len(str(values[column - 1])) + 2))
        sheet.column_dimensions[get_column_letter(column)].width = width
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.oddFooter.right.text = "Página &P"


def stamp_logo(sheet: Worksheet, title: str, subtitle: str) -> None:
    image = SheetImage(str(LOGO))
    image.width = 210
    image.height = 60
    sheet.add_image(image, "A1")
    sheet.merge_cells("D1:G1")
    sheet.merge_cells("D2:G2")
    title_cell = sheet["D1"]
    title_cell.value = title
    title_cell.font = Font(name="Calibri", bold=True, size=18, color=NAVY_HEX)
    title_cell.alignment = Alignment(vertical="center")
    subtitle_cell = sheet["D2"]
    subtitle_cell.value = subtitle
    subtitle_cell.font = Font(name="Calibri", size=12, color="5A6A80")
    subtitle_cell.alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 24
    sheet.row_dimensions[2].height = 20
    sheet.row_dimensions[3].height = 12
    sheet.row_dimensions[4].height = 8
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = "landscape"
    sheet.print_title_rows = "1:5"
    sheet.page_margins.top = 0.6
    sheet.page_margins.header = 0.2


def workbook_bytes(book: Workbook) -> bytes:
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def sum_rows(parts: list[PartSum]) -> tuple[list[str], list[list]]:
    headers = ["Serie", "Descripción", "Tipo", "Faltante", "Incompleto", "Inoperativo", "Total"]
    rows = [
        [part.code, part.description, part.part_type or "", part.faltante, part.incompleto, part.inoperativo, part.total]
        for part in parts
    ]
    return headers, rows


def build_daily_xlsx(report: DailyReport) -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.title = "Inspecciones"
    stamp_logo(sheet, "Máquinas inspeccionadas", format_day(report.date))
    headers = ["Técnico", "Serie", "Marca", "Modelo", "Ubicación", "Repuestos", "Tiempo", "Cierre"]
    rows = [
        [
            group.technician,
            row.serial,
            row.brand,
            row.model,
            row.location or "",
            row.items_count,
            format_duration(row.assigned_at, row.completed_at),
            format_when(row.completed_at),
        ]
        for group in report.groups
        for row in group.inspections
    ]
    write_table(sheet, 5, headers, rows)
    return workbook_bytes(book)


def excel_sheet_name(title: str, used: set[str]) -> str:
    cleaned = " ".join("".join(" " if char in "[]:*?/\\" else char for char in title).split()) or "Modelo"
    base = cleaned[:31]
    name = base
    number = 2
    while name in used:
        suffix = f" {number}"
        name = f"{base[: 31 - len(suffix)]}{suffix}"
        number += 1
    used.add(name)
    return name


def unit_rows(units: list) -> list[list]:
    rows = []
    for unit in units:
        if not unit.parts:
            rows.append([
                unit.serial, unit.location or "", unit.year or "", unit.technician,
                format_when(unit.completed_at), "", "", "", "", "Sin faltantes",
            ])
            continue
        for part in unit.parts:
            rows.append([
                unit.serial,
                unit.location or "",
                unit.year or "",
                unit.technician,
                format_when(unit.completed_at),
                part.quantity,
                CONDITIONS.get(part.condition, part.condition),
                part.code,
                part.description,
                part.note or "",
            ])
    return rows


def build_needs_xlsx(report: NeedsReport, view: str) -> bytes:
    book = Workbook()
    sheet = book.active
    detail_headers = ["Serie equipo", "Ubicación", "Año", "Técnico", "Cierre", "Cant.", "Condición", "Serie", "Descripción", "Nota"]
    if view == "general":
        span = f"Del {format_day(report.date_from)} al {format_day(report.date_to)}"
        if not report.models:
            sheet.title = "General"
            stamp_logo(sheet, "Reporte general", span)
            write_table(sheet, 5, detail_headers, [])
            return workbook_bytes(book)
        used: set[str] = set()
        for index, block in enumerate(report.models):
            pieces = sum(part.total for part in block.parts)
            current = sheet if index == 0 else book.create_sheet()
            current.title = excel_sheet_name(f"{block.brand} {block.model}", used)
            stamp_logo(
                current,
                f"{block.brand} {block.model}",
                f"{span}  ·  {block.machines} máquinas  ·  {pieces} piezas",
            )
            write_table(current, 5, detail_headers, unit_rows(block.units))
        return workbook_bytes(book)
    if view == "modelo":
        block = report.by_model
        pieces = sum(part.total for part in block.parts)
        sheet.title = "Por modelo"
        stamp_logo(
            sheet,
            f"Necesidad de {block.brand} {block.model}",
            f"Del {format_day(report.date_from)} al {format_day(report.date_to)}  ·  {block.machines} máquinas  ·  {pieces} piezas",
        )
        write_table(sheet, 5, detail_headers, unit_rows(block.units))
        total = book.create_sheet("Total del modelo")
        stamp_logo(total, f"Total de {block.brand} {block.model}", f"{block.machines} equipos  ·  {pieces} piezas")
        total_headers, total_rows = sum_rows(block.parts)
        write_table(total, 5, total_headers, total_rows)
    elif view == "total":
        sheet.title = "Conglomerado"
        stamp_logo(
            sheet,
            "Conglomerado de repuestos",
            f"Del {format_day(report.date_from)} al {format_day(report.date_to)}  ·  {report.summary.inspections} inspecciones  ·  {report.summary.parts_total} piezas",
        )
        headers, rows = sum_rows(report.conglomerate)
        write_table(sheet, 5, headers, rows)
    else:
        sheet.title = "Por serie"
        stamp_logo(
            sheet,
            "Necesidad por serie",
            f"Del {format_day(report.date_from)} al {format_day(report.date_to)}  ·  {report.summary.inspections} inspecciones  ·  {report.summary.parts_total} piezas",
        )
        headers = ["Serie máquina", "Marca", "Modelo", "Ubicación", "Año", "Técnico", "Cierre", "Cant.", "Condición", "Serie", "Descripción", "Nota"]
        rows = []
        for block in report.by_serial:
            if not block.parts:
                rows.append([
                    block.serial, block.brand, block.model, block.location or "", block.year or "",
                    block.technician, format_when(block.completed_at), "", "", "", "", "Sin faltantes",
                ])
                continue
            for part in block.parts:
                rows.append([
                    block.serial,
                    block.brand,
                    block.model,
                    block.location or "",
                    block.year or "",
                    block.technician,
                    format_when(block.completed_at),
                    part.quantity,
                    CONDITIONS.get(part.condition, part.condition),
                    part.code,
                    part.description,
                    part.note or "",
                ])
        write_table(sheet, 5, headers, rows)
    return workbook_bytes(book)
